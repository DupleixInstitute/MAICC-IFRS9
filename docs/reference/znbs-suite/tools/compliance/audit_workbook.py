"""Directive compliance audit - one engine for every BoZ directive.

Each directive is a data module in tools/compliance/directives/ exposing META, ROWS and
FINDINGS. This module turns that data into the three deliverables:

  * an Excel workbook: 'Contents' in the directive's own arrangement of sections, each
    section hyperlinked to its row on 'Audit'; 'Audit' with the nine audit columns, a
    status drop-down and colour coding; 'Findings' for what needs a decision or a fix;
  * a Markdown twin generated from the same rows (so the two cannot drift);
  * a PDF of the Markdown (tools/compliance/md_to_pdf.py).

ROWS:     (part, ref, name, action, status, tool_comment, general_comment,
           app_compliance_comment, app_reports_reference)
FINDINGS: (no, directive_ref, finding, what_was_found, impact, recommended_action,
           owner, status)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.formatting.rule import FormulaRule

DONE = "Done"
PART = "Partially Done"
OUT_ = "Outstanding"
NA = "Not Applicable - documented"
EVID = "Evidence Needed - ZNBS"
STATUSES = [DONE, PART, OUT_, NA, EVID]
STATUS_FILL = {DONE: "C6EFCE", PART: "FFEB9C", OUT_: "FFC7CE", NA: "D9D9D9", EVID: "DDEBF7"}
STATUS_DEFS = {
    DONE: "Implemented in the tool and visible in an approved output (run, return or pack).",
    PART: "The mechanics exist, but a parameter, an input or a piece of evidence differs from what the directive requires.",
    OUT_: "Required for ZNBS and not implemented.",
    NA: "Does not apply to ZNBS by a recorded decision or fact (reason and reference given); revisit if the facts change.",
    EVID: "The tool is ready; ZNBS must supply a document, minute or dataset before the row can be closed.",
}
CONTEXT = "Zambia National Building Society - ICAAP as at 31 December 2025 - final checks (UAT)."

FONT = "Arial"
GOLD = "B8860B"; PRIMARY = "1F3B2D"; CREAM = "FBF0E4"; GREY = "595959"; LINK = "0563C1"
_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

HEAD = ["Directive section ref", "Directive section name", "Directive action (what the section requires)",
        "Tool's status", "Tool comment (what the tool does / evidence checked)", "General comment",
        "App compliance comment", "App reports reference (where to see it)",
        "ZNBS reviewer sign-off (initials / date)"]
WIDTHS = [12, 34, 60, 22, 60, 42, 44, 52, 22]


def _link(cell, sheet, target):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{target}", display=str(cell.value))
    cell.font = Font(name=FONT, size=9, color=LINK, underline="single")


def _header(ws, row, labels, widths):
    for i, (h, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=PRIMARY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w


def validate(meta, rows, findings):
    """Fail loudly on malformed data rather than write a misleading workbook."""
    refs = set()
    for r in rows:
        assert len(r) == 9, f"row needs 9 fields: {r[:3]}"
        assert r[4] in STATUSES, f"unknown status {r[4]!r} on {r[1]}"
        assert r[1] not in refs, f"duplicate ref {r[1]}"
        refs.add(r[1])
        if r[4] == NA:
            assert r[5].strip() or r[7].strip(), f"N/A without a reason on {r[1]}"
    for f in findings:
        assert len(f) == 8, f"finding needs 8 fields: {f[:2]}"
    for k in ("file_stem", "short", "title", "basis", "source"):
        assert meta.get(k), f"META.{k} missing"


def build_xlsx(meta, rows, findings, out):
    validate(meta, rows, findings)
    wb = Workbook()

    # Audit
    ws = wb.active
    ws.title = "Audit"
    ws["A1"] = meta["title"] + " - compliance audit"
    ws["A1"].font = Font(name=FONT, size=13, bold=True, color=PRIMARY)
    ws["A2"] = f"{CONTEXT} {meta['basis']} Status definitions: see 'Contents'."
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color=GREY)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 30
    back = ws.cell(row=3, column=1, value="Back to Contents")
    _link(back, "Contents", "A1")
    hr = 4
    _header(ws, hr, HEAD, WIDTHS)
    ws.row_dimensions[hr].height = 34
    row = hr + 1
    part_rows, ref_rows = {}, {}
    for part, ref, name, action, status, tool, general, comp, appref in rows:
        if part not in part_rows:
            ws.cell(row=row, column=1, value=part).font = Font(name=FONT, bold=True, size=10, color=PRIMARY)
            for col in range(1, 10):
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=CREAM)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
            part_rows[part] = row
            row += 1
        for col, v in enumerate([ref, name, action, status, tool, general, comp, appref, ""], start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = Font(name=FONT, size=9, bold=(col == 4))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
        ref_rows[ref] = row
        row += 1
    last = row - 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A{hr}:I{last}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D{hr + 1}:D{last}")
    for st, colr in STATUS_FILL.items():
        ws.conditional_formatting.add(f"D{hr + 1}:D{last}", FormulaRule(formula=[f'$D{hr + 1}="{st}"'], fill=PatternFill("solid", fgColor=colr)))
    ws.sheet_properties.tabColor = PRIMARY

    # Contents
    toc = wb.create_sheet("Contents", 0)
    toc.sheet_properties.tabColor = GOLD
    toc["A1"] = meta["short"] + " - compliance audit"
    toc["A1"].font = Font(name=FONT, size=14, bold=True, color=PRIMARY)
    toc["A2"] = "Contents (the directive's own arrangement of sections). Click a section to open its row on the Audit sheet."
    toc["A2"].font = Font(name=FONT, size=9, italic=True, color=GREY)
    toc["A4"] = "Status definitions"
    toc["A4"].font = Font(name=FONT, bold=True, size=10, color=PRIMARY)
    toc["C4"] = "Sections"
    toc["C4"].font = Font(name=FONT, bold=True, size=10, color=PRIMARY)
    toc["C4"].alignment = Alignment(horizontal="center")
    r = 5
    for st in STATUSES:
        c = toc.cell(row=r, column=1, value=st)
        c.font = Font(name=FONT, size=9, bold=True)
        c.fill = PatternFill("solid", fgColor=STATUS_FILL[st])
        c.border = BORDER
        c = toc.cell(row=r, column=2, value=STATUS_DEFS[st])
        c.font = Font(name=FONT, size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c = toc.cell(row=r, column=3, value=f'=COUNTIF(Audit!$D:$D,"{st}")')
        c.font = Font(name=FONT, size=9, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="top")
        toc.row_dimensions[r].height = 26
        r += 1
    toc.cell(row=r, column=1, value="Sections audited").font = Font(name=FONT, size=9, bold=True)
    c = toc.cell(row=r, column=3, value=f"=COUNTA(Audit!$A$5:$A${last})-{len(part_rows)}")
    c.font = Font(name=FONT, size=9, bold=True)
    c.alignment = Alignment(horizontal="center")
    r += 2
    _header(toc, r, ["Ref", "Section", "Tool's status", "Open in Audit"], [30, 62, 30, 16])
    first_toc = r + 1
    r += 1
    seen = set()
    for part, ref, name, *_ in rows:
        if part not in seen:
            seen.add(part)
            toc.cell(row=r, column=1, value=part).font = Font(name=FONT, bold=True, size=10, color=PRIMARY)
            for col in range(1, 5):
                toc.cell(row=r, column=col).fill = PatternFill("solid", fgColor=CREAM)
            toc.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            r += 1
        ar = ref_rows[ref]
        c = toc.cell(row=r, column=1, value=ref); c.font = Font(name=FONT, size=9); c.border = BORDER
        c = toc.cell(row=r, column=2, value=name); _link(c, "Audit", f"A{ar}"); c.border = BORDER
        c = toc.cell(row=r, column=3, value=f"=Audit!D{ar}"); c.font = Font(name=FONT, size=9); c.border = BORDER
        c = toc.cell(row=r, column=4, value=f"Audit row {ar}"); _link(c, "Audit", f"A{ar}"); c.border = BORDER
        r += 1
    for st, colr in STATUS_FILL.items():
        toc.conditional_formatting.add(f"C{first_toc}:C{r}", FormulaRule(formula=[f'$C{first_toc}="{st}"'], fill=PatternFill("solid", fgColor=colr)))
    toc.freeze_panes = "A3"
    r += 1
    c = toc.cell(row=r, column=1, value="Findings that need a decision or a fix: see the 'Findings' sheet")
    _link(c, "Findings", "A1")
    r += 1
    c = toc.cell(row=r, column=1, value="Source: " + meta["source"])
    c.font = Font(name=FONT, size=8, italic=True, color=GREY)

    # Findings
    fs = wb.create_sheet("Findings")
    fs.sheet_properties.tabColor = "C00000"
    fs["A1"] = "Findings - items that need a decision or a fix before this directive can be signed off"
    fs["A1"].font = Font(name=FONT, size=13, bold=True, color=PRIMARY)
    _link(fs.cell(row=2, column=1, value="Back to Contents"), "Contents", "A1")
    _header(fs, 4, ["No.", "Directive ref", "Finding", "What was found", "Impact at Dec-2025", "Recommended action", "Owner", "Status"],
            [6, 16, 30, 64, 40, 64, 26, 12])
    for j, f in enumerate(findings, start=5):
        for i, v in enumerate(f, start=1):
            c = fs.cell(row=j, column=i, value=v)
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
            if i == 2:
                first = str(v).replace("s.", "").split(",")[0].split("(")[0].split(" ")[0].strip()
                if first in ref_rows:
                    _link(c, "Audit", f"A{ref_rows[first]}")
    fs.freeze_panes = "A5"
    if findings:
        dv2 = DataValidation(type="list", formula1='"Open,In progress,Closed,Deferred"', allow_blank=True)
        fs.add_data_validation(dv2)
        dv2.add(f"H5:H{4 + len(findings)}")

    for sh in (toc, ws, fs):
        sh.sheet_view.showGridLines = False
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    fs.print_title_rows = "4:4"
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return ref_rows


def build_markdown(meta, rows, findings, out):
    counts = {st: sum(1 for r in rows if r[4] == st) for st in STATUSES}
    L = [f"# {meta['short']} - compliance audit", "",
         f"{meta['title']}.", "",
         f"{CONTEXT} {meta['basis']} The Excel workbook of the same name is the working copy "
         "(drop-down statuses, hyperlinked contents, reviewer sign-off column); this document is its printable twin.", "",
         f"Source: {meta['source']}", "",
         "## Status summary", "",
         "| Status | Meaning | Sections |", "|---|---|---|"]
    for st in STATUSES:
        L.append(f"| **{st}** | {STATUS_DEFS[st]} | {counts[st]} |")
    L += [f"| **Total** | | {len(rows)} |", "", "## Contents", "", "| Ref | Section | Tool's status |", "|---|---|---|"]
    part = None
    for p, ref, name, action, status, *_ in rows:
        if p != part:
            L.append(f"| **{p}** | | |")
            part = p
        L.append(f"| {ref} | {name} | {status} |")
    L += ["", "## Findings - items that need a decision or a fix", ""]
    if not findings:
        L += ["No findings.", ""]
    for no, ref, title, found, impact, action, owner, status in findings:
        L += [f"### {no}. {title} (s.{ref}) - {status}", "",
              f"- **What was found:** {found}", f"- **Impact at Dec-2025:** {impact}",
              f"- **Recommended action:** {action}", f"- **Owner:** {owner}", ""]
    L.append("## Section-by-section audit")
    part = None
    for p, ref, name, action, status, tool, general, comp, appref in rows:
        if p != part:
            L += ["", f"### {p}"]
            part = p
        L += ["", f"**{ref}. {name}** - *{status}*", "",
              f"- **Directive action:** {action}", f"- **Tool comment:** {tool}"]
        if general:
            L.append(f"- **General comment:** {general}")
        L += [f"- **App compliance comment:** {comp}", f"- **App reports reference:** {appref}",
              "- **ZNBS reviewer sign-off:** ____________________   Date: ____________"]
    L.append("")
    with open(out, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(L))
