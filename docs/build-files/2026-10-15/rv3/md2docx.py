"""Markdown to Word for the results review: headings, paragraphs, bullets, numbered items,
pipe tables, **bold** and `code` runs; logo, running header, page numbers.
usage: python -I md2docx.py in.md out.docx logo.png"""
import sys, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out, logo = sys.argv[1], sys.argv[2], sys.argv[3]
lines = open(src, encoding='utf-8').read().splitlines()
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xF0, 0x8A, 0x1C)
GREY = RGBColor(0x55, 0x5B, 0x66)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.7)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)
sec.header_distance = sec.footer_distance = Cm(0.7)
sec.different_first_page_header_footer = True
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.08
for lvl, size, before in ((1, 15, 14), (2, 12, 10), (3, 11, 8)):
    h = doc.styles[f'Heading {lvl}']
    h.font.name = 'Calibri'
    h.element.rPr.rFonts.set(qn('w:asciiTheme'), 'minorHAnsi') if h.element.rPr is not None else None
    h.font.size = Pt(size)
    h.font.color.rgb = NAVY
    h.font.bold = True
    h.paragraph_format.space_before = Pt(before)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
for name in ('List Bullet', 'List Number'):
    doc.styles[name].font.size = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(3)


def field(run, code):
    for kind, text in (('begin', None), (None, code), ('separate', None), (None, '1'), ('end', None)):
        if kind:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind); run._r.append(e)
        elif text == code:
            t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = code; run._r.append(t)
        else:
            t = OxmlElement('w:t'); t.text = text; run._r.append(t)


title = lines[0][2:] if lines and lines[0].startswith('# ') else ''
# running header and footer (not on the first page)
hp = sec.header.paragraphs[0]
hr = hp.add_run(title + '   |   Dupleix Institute   |   Draft for Edward Mazibuko\'s review')
hr.font.size = Pt(8); hr.font.color.rgb = GREY
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = fp.add_run('Page '); r.font.size = Pt(8); r.font.color.rgb = GREY
r = fp.add_run(); r.font.size = Pt(8); r.font.color.rgb = GREY; field(r, 'PAGE')
r = fp.add_run(' of '); r.font.size = Pt(8); r.font.color.rgb = GREY
r = fp.add_run(); r.font.size = Pt(8); r.font.color.rgb = GREY; field(r, 'NUMPAGES')
ffp = sec.first_page_footer.paragraphs[0]
ffp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = ffp.add_run('Dupleix Institute   |   Confidential: MAIIC client data'); r.font.size = Pt(8); r.font.color.rgb = GREY


def runs(p, text, size=None, color=None, bold=None):
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt((size or 10) - 1)
        else:
            r = p.add_run(part)
        if size and not (part.startswith('`')):
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        if bold:
            r.bold = True


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)


def borders(table, color='B8C2CF'):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for k in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color); b.append(e)
    tblPr.append(b)


NUM = re.compile(r'^[-+(]?[0-9][0-9.,]*%?\)?( \([0-9,]+\))?( / [-+]?[0-9][0-9.,]*%?)*$|^-$')

# cover block
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.add_run().add_picture(logo, width=Cm(5.2))
i = 0
if lines and lines[0].startswith('# '):
    p = doc.add_paragraph()
    r = p.add_run(title); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    r = p.add_run('Expected credit loss and effective interest rate: the IFRS 9 system against MAIIC\'s December 2025 models, the signed 2025 statements and MAIIC\'s manual workings')
    r.font.size = Pt(11.5); r.font.color.rgb = ORANGE
    p.paragraph_format.space_after = Pt(10)
    i = 1

while i < len(lines):
    ln = lines[i]
    if ln.startswith('|'):
        tbl = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                tbl.append(cells)
            i += 1
        ncol = max(len(r) for r in tbl)
        t = doc.add_table(rows=len(tbl), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True
        borders(t)
        cell_margins(t)
        size = 8 if ncol >= 9 else 8.5
        numcol = {}
        for ci in range(1, ncol):
            vals = [r_[ci].replace('**', '') for r_ in tbl[1:] if ci < len(r_) and r_[ci].strip()]
            numcol[ci] = bool(vals) and sum(1 for x in vals if NUM.match(x) or re.match(r'^[-+]?[0-9][0-9,.]*%? (on|[(]|against)', x)) >= len(vals) * 0.6
        for ri, row in enumerate(tbl):
            tr = t.rows[ri]._tr
            trPr = tr.get_or_add_trPr()
            cs = OxmlElement('w:cantSplit'); trPr.append(cs)
            if ri == 0:
                th = OxmlElement('w:tblHeader'); trPr.append(th)
            for ci in range(ncol):
                cell = t.cell(ri, ci)
                cell.text = ''
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                txt = row[ci] if ci < len(row) else ''
                if ri == 0:
                    runs(p, txt, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                    shade(cell, '1F3A5F')
                else:
                    runs(p, txt, size=size)
                    if ri % 2 == 0:
                        shade(cell, 'F2F5F9')
                if ci > 0 and numcol.get(ci):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2)
        continue
    if ln.startswith('# '):
        pass
    elif ln.startswith('## '):
        h = doc.add_heading(ln[3:], 1)
        if ln[3:].startswith('Appendix A'):
            h.paragraph_format.page_break_before = True
    elif ln.startswith('### '):
        doc.add_heading(ln[4:], 2)
    elif re.match(r'^\s*- ', ln):
        p = doc.add_paragraph(style='List Bullet'); runs(p, re.sub(r'^\s*- ', '', ln))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif re.match(r'^\d+\. ', ln):
        p = doc.add_paragraph(style='List Number'); runs(p, re.sub(r'^\d+\. ', '', ln))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif ln.strip() == '':
        pass
    else:
        p = doc.add_paragraph(); runs(p, ln)
        if re.match(r'^\*\*(Table \d+|[A-Z][^*]{0,80}(window|cohort)[^*]*)\.?', ln) and len(ln) < 260:
            p.paragraph_format.keep_with_next = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    i += 1
doc.core_properties.title = title
doc.core_properties.author = 'Dupleix Institute'
doc.save(out)
print('saved', out)
