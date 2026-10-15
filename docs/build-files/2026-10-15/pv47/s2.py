import openpyxl, sys, re, warnings
warnings.filterwarnings('ignore')
from openpyxl.utils import get_column_letter as L
p=sys.argv[1]
wf=openpyxl.load_workbook(p,data_only=False)
wv=openpyxl.load_workbook(p,data_only=True)
print("== refs to CB7")
pat=re.compile(r"\$?CB\$?7(?!\d)")
for ws in wf.worksheets:
  for row in ws.iter_rows():
    for c in row:
      if isinstance(c.value,str) and c.value.startswith('=') and pat.search(c.value):
        pass
cnt={}
for ws in wf.worksheets:
  for row in ws.iter_rows():
    for c in row:
      if isinstance(c.value,str) and c.value.startswith('=') and pat.search(c.value):
        k=(ws.title,c.column_letter); cnt.setdefault(k,[]).append((c.row,c.value))
for k,v in cnt.items(): print(k,len(v),v[:2],v[-1])
print("== ECLComp headers rows 1-10 all cols")
f=wf['ECLComp']; v=wv['ECLComp']
for c in range(1,f.max_column+1):
  hs=[str(f.cell(r,c).value).replace('\n',' ') for r in range(1,11) if f.cell(r,c).value is not None]
  print(L(c), '|', ' / '.join(hs)[:160], '|R11:', str(f.cell(11,c).value)[:200], '=>', v.cell(11,c).value)
