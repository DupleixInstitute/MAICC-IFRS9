import openpyxl, sys, pickle
p=sys.argv[1]
wf=openpyxl.load_workbook(p,data_only=False)
wv=openpyxl.load_workbook(p,data_only=True)
print(wf.sheetnames)
for ws in wf.worksheets: print(ws.title, ws.dimensions)
print("NAMES")
for n,d in wf.defined_names.items(): print(n, d.attr_text)
for ws in wf.worksheets:
    for n,d in ws.defined_names.items(): print(ws.title, n, d.attr_text)
f=wf['ECLComp']; v=wv['ECLComp']
for r in range(1,12):
  for c in range(openpyxl.utils.column_index_from_string('AZ'),openpyxl.utils.column_index_from_string('CE')+1):
    cf=f.cell(r,c).value
    if cf is not None: print(f.cell(r,c).coordinate, repr(cf), '=>', v.cell(r,c).value)
