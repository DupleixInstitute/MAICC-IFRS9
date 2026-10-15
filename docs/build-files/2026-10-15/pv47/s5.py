import openpyxl, sys, re, warnings
warnings.filterwarnings('ignore')
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L
p=sys.argv[1]
wf=openpyxl.load_workbook(p,data_only=False)
wv=openpyxl.load_workbook(p,data_only=True)
def dump(s,r1,r2,c1,c2,w=80):
  f=wf[s]; v=wv[s]
  for r in range(r1,r2+1):
    out=[]
    for c in range(CI(c1),CI(c2)+1):
      x=f.cell(r,c).value
      if x is not None: out.append(f"{L(c)}{r}={str(x)[:w]!r}->{str(v.cell(r,c).value)[:20]}")
    if out: print(' | '.join(out))
print('comment CB7:', wf['ECLComp']['CB7'].comment, wf['ECLComp']['CA7'].comment, wf['ECLComp']['CA5'].comment)
for c in ['CA4','CB4','CA5','CB5','CA6','CB6','CA7','CB7','CC7','CA8','CB8','CC5','CC6']:
  print(c, repr(wf['ECLComp'][c].value), wf['ECLComp'][c].number_format, wf['ECLComp'][c].fill.fgColor.rgb)
print('== NPL sheet BK210-220 area'); dump('NPL',205,222,'BE','BM')
print('== NPL B180-186'); dump('NPL',180,186,'A','H')
print('== GDP U80-86'); dump('GDP',80,86,'P','W')
print('== Methodology FLI rows 8-20'); dump('Methodology Foward looking(FLI)',8,20,'A','F',120)
print('== KPIs'); dump('KPIs',2,13,'B','G',60)
print('== Summary(OriginalMethod) 6-16'); dump('Summary(OriginalMethod)',6,16,'B','O',60)
