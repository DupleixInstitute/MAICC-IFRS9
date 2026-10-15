import openpyxl, sys, re, warnings
warnings.filterwarnings('ignore')
p=sys.argv[1]
wf=openpyxl.load_workbook(p,data_only=False)
wv=openpyxl.load_workbook(p,data_only=True)
kw=re.compile(r'recover|\bNPL|cure|from npl|46\.98|0\.4698|0\.46979|write.?off|LGD|loss given',re.I)
print('== text hits')
for ws in wf.worksheets:
  hits=[]
  for row in ws.iter_rows():
    for c in row:
      val=c.value
      if isinstance(val,str) and kw.search(val):
        hits.append((c.coordinate,val[:110].replace('\n',' ')))
  if hits: print(ws.title,len(hits),hits[:25])
print('== numeric near 0.4698 / 46.98')
for ws in wv.worksheets:
  for row in ws.iter_rows():
    for c in row:
      x=c.value
      if isinstance(x,(int,float)) and not isinstance(x,bool) and (abs(x-0.4697984)<5e-4 or abs(x-46.97984)<5e-2 or abs(x-0.5302016)<5e-4 or abs(x-53.02016)<5e-2):
        print(ws.title,c.coordinate,x,wf[ws.title][c.coordinate].value)
print('== sister sheets col CF and their ECL formula')
for s in ['ECLComp(InternalPDByCount) ','ECLComp(InternalPDByValue)']:
  f=wf[s]; v=wv[s]
  for r in range(4,12):
    for col in ['CE','CF','CG','BF','BG','BE','BI']:
      if f[f'{col}{r}'].value is not None: print(s,col,r,repr(f[f'{col}{r}'].value)[:150],'=>',v[f'{col}{r}'].value)
