import openpyxl, sys, warnings, bisect
warnings.filterwarnings('ignore')
p=sys.argv[1]
wv=openpyxl.load_workbook(p,data_only=True,read_only=True)
target=0.4697984013775794
pool={}
for s in wv.sheetnames:
  if s.startswith('ECLComp'): continue
  for row in wv[s].iter_rows():
    for c in row:
      x=getattr(c,'value',None)
      if isinstance(x,(int,float)) and not isinstance(x,bool) and abs(x)>1:
        pool.setdefault(round(abs(float(x)),6),(s,c.coordinate))
vals=sorted(pool)
print('distinct',len(vals))
hits=[]
for a in vals:
  for r,tag in ((target,'r'),(1-target,'1-r')):
    t=a*r; i=bisect.bisect_left(vals,t*(1-2e-8))
    while i<len(vals) and vals[i]<=t*(1+2e-8):
      hits.append((tag,vals[i],pool[vals[i]],a,pool[a])); i+=1
print('hits',len(hits))
for h in hits[:30]: print(h)
