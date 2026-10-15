"""Minimal Markdown to Word converter for the review report: # headings, paragraphs,
- bullets, 1. numbered items, pipe tables, **bold** runs. usage: md2docx.py in.md out.docx"""
import sys, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out = sys.argv[1], sys.argv[2]
lines = open(src, encoding='utf-8').read().splitlines()
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
for lvl, size in ((1, 16), (2, 13), (3, 11)):
    h = doc.styles[f'Heading {lvl}']
    h.font.name = 'Calibri'
    h.font.size = Pt(size)
    h.font.color.rgb = NAVY


def runs(p, text):
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else:
            p.add_run(part)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith('|'):
        tbl = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                tbl.append(cells)
            i += 1
        ncol = max(len(r) for r in tbl)
        t = doc.add_table(rows=len(tbl), cols=ncol)
        t.style = 'Table Grid'
        for ri, r in enumerate(tbl):
            for ci in range(ncol):
                cell = t.cell(ri, ci)
                cell.text = ''
                p = cell.paragraphs[0]
                txt = r[ci] if ci < len(r) else ''
                runs(p, txt)
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    if ri == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                if ri == 0:
                    shade(cell, '1F3A5F')
                if ci > 0 and re.fullmatch(r'[-()0-9.,% ]+[a-z]?', txt or 'x'):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()
        continue
    if ln.startswith('# '):
        doc.add_heading(ln[2:], 0 if i == 0 else 1)
    elif ln.startswith('## '):
        doc.add_heading(ln[3:], 1)
    elif ln.startswith('### '):
        doc.add_heading(ln[4:], 2)
    elif re.match(r'^\s*- ', ln):
        p = doc.add_paragraph(style='List Bullet'); runs(p, re.sub(r'^\s*- ', '', ln))
    elif re.match(r'^\d+\. ', ln):
        p = doc.add_paragraph(style='List Number'); runs(p, re.sub(r'^\d+\. ', '', ln))
    elif ln.strip() == '':
        pass
    else:
        p = doc.add_paragraph(); runs(p, ln)
    i += 1
doc.save(out)
print('saved', out)
