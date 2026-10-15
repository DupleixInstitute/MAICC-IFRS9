import sys, openpyxl, csv
from openpyxl.utils import column_index_from_string as ci
wb = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
ws = wb['ECLComp']
cols=['A','B','E','F','H','I','P','R','AR','AS','AT','AU','AV','AW','BB','BC']
out=csv.writer(open(sys.argv[2],'w',newline='',encoding='utf-8'))
out.writerow(cols)
for row in ws.iter_rows(min_row=11,max_row=ws.max_row,values_only=True):
    if row[1] in (None,''): continue
    out.writerow([row[ci(c)-1] for c in cols])
