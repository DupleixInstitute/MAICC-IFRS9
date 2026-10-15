import sys, os, csv, io
import openpyxl, xlrd

ROOT = sys.argv[1]
SKIP = ("Received 9 October", "Received 12 October 2026", "Build files")
EXT = (".xlsx", ".xls", ".csv", ".txt")
only = sys.argv[2] if len(sys.argv) > 2 else None


def short(v):
    s = str(v).strip().replace("\n", " ")
    return s[:28]


def first_rows_xlsx(path):
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        try:
            dims = ws.calculate_dimension()
        except Exception:
            dims = "?"
        rows = []
        n = 0
        for r in ws.iter_rows(values_only=True):
            n += 1
            if len(rows) < 4 and any(c not in (None, "") for c in r):
                rows.append([short(c) for c in r if c not in (None, "")][:14])
            if n > 200000:
                break
        out.append((ws.title, dims, ws.max_row, rows))
    wb.close()
    return out


def first_rows_xls(path):
    out = []
    try:
        wb = xlrd.open_workbook(path, on_demand=True)
    except Exception as e:
        # maybe html/xml disguised
        with open(path, "rb") as f:
            head = f.read(300)
        return [("ERR", str(e)[:80], 0, [[head[:120]]])]
    for sh in wb.sheets():
        rows = []
        for i in range(min(sh.nrows, 60)):
            r = sh.row_values(i)
            if any(c not in (None, "") for c in r) and len(rows) < 4:
                rows.append([short(c) for c in r if c not in (None, "")][:14])
        out.append((sh.name, f"{sh.nrows}x{sh.ncols}", sh.nrows, rows))
    return out


def csv_info(path):
    with open(path, "rb") as f:
        data = f.read()
    txt = data.decode("utf-8", "replace")
    lines = txt.splitlines()
    return [("csv", "", len(lines), [lines[0][:300], (lines[1][:200] if len(lines) > 1 else ""), (lines[-1][:200] if lines else "")])]


for dp, dn, fn in os.walk(ROOT):
    if any(s in dp for s in SKIP):
        continue
    for f in sorted(fn):
        p = os.path.join(dp, f)
        rel = os.path.relpath(p, ROOT)
        if only and only not in rel:
            continue
        low = f.lower()
        if not low.endswith(EXT):
            continue
        try:
            if low.endswith(".xlsx"):
                info = first_rows_xlsx(p)
            elif low.endswith(".xls"):
                info = first_rows_xls(p)
            else:
                info = csv_info(p)
        except Exception as e:
            info = [("ERR", str(e)[:120], 0, [])]
        print("=" * 4, rel)
        for name, dims, n, rows in info[:25]:
            print(f"  [{name}] {dims} rows={n}")
            for r in rows:
                print("     ", r)
        if len(info) > 25:
            print(f"  ... {len(info)} sheets total")
