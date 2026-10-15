"""Reconcile Finance's 2025 fee workbook (per loan) to the E-Banker GL fee postings on 4871/4873 (read-only)."""
import subprocess, sys, re
from datetime import datetime, date
from collections import defaultdict
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FIN, OUT = sys.argv[1], sys.argv[2]
M = r"C:\xampp\mysql\bin\mysql.exe"


def rows(sql):
    r = subprocess.run([M, "-uroot", "-B", "--default-character-set=utf8mb4", "maiic_ifrs9"], input=sql,
                       capture_output=True, text=True, encoding="utf-8")
    if r.returncode:
        raise SystemExit(r.stderr)
    lines = r.stdout.rstrip("\n").split("\n")
    head = lines[0].split("\t")
    return [dict(zip(head, [None if v == "NULL" else v for v in l.split("\t")])) for l in lines[1:]]


# ---------------------------------------------------------------- Finance's workbook
ws = openpyxl.load_workbook(FIN, data_only=True).active
fin = []
for i, r in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
    acct = r[3]
    if not acct or not str(acct).strip().isdigit():
        continue
    vd = r[7]
    if isinstance(vd, str):
        vd = datetime.strptime(vd.strip(), "%d/%m/%Y")
    fin.append({"row": i, "acct": str(acct).strip().lstrip("0"), "name": (r[5] or r[4] or "").strip(), "value_date": vd.date() if vd else None,
                "legal": float(r[18] or 0), "arrangement": float(r[20] or 0)})
fin_tot = {"legal": sum(f["legal"] for f in fin), "arrangement": sum(f["arrangement"] for f in fin)}

# ---------------------------------------------------------------- ledger postings, 2025, GL 4871 legal / 4873 arrangement
led = rows("""
SELECT CONCAT('T', id) pid, 'Tied to a loan by the system' state, contract_id loan, gl_account_ref gl, amount, transaction_date d, description narration, source_reference ref
FROM contract_fees WHERE source_system = 'E-Banker GL fee posting' AND gl_account_ref IN ('4871','4873') AND YEAR(transaction_date) = 2025
UNION ALL
SELECT CONCAT('H', id), 'Held: no loan in the batch', candidate_contract_id, gl_code, amount, transaction_date, narration, CONCAT(batch, ' / ', external_transaction_id)
FROM fee_gl_unassigned WHERE gl_code IN ('4871','4873') AND YEAR(transaction_date) = 2025""")
for p in led:
    p["amount"] = float(p["amount"]); p["d"] = date.fromisoformat(p["d"][:10]); p["type"] = "legal" if p["gl"] == "4871" else "arrangement"
    p["used"] = False
led_tot = {t: sum(p["amount"] for p in led if p["type"] == t) for t in ("legal", "arrangement")}
names = {f["acct"]: f["name"] for f in fin}


def name_hit(p, f):
    words = [w for w in re.findall(r"[a-z]{4,}", f["name"].lower()) if w not in ("limited", "farms", "farm", "sacco", "services", "investment", "investments", "holdings", "financial", "general", "company")]
    n = (p["narration"] or "").lower()
    return any(w in n for w in words)


# ---------------------------------------------------------------- match
from itertools import combinations

items = [{"f": f, "t": t, "want": round(f[t], 2), "posts": [], "how": []} for f in fin for t in ("legal", "arrangement") if round(f[t], 2)]


def days(p, f):
    return abs((p["d"] - f["value_date"]).days) if f["value_date"] else 0


def take(it, ps, why):
    for p in ps:
        p["used"] = True
        it["posts"].append(p)
        w = why
        if p["state"].startswith("Tied") and p["loan"] != it["f"]["acct"]:
            w = f"the system tied it to loan {p['loan']} (that loan is the leg of the batch); Finance's file puts it on this loan"
        elif p["state"].startswith("Held") and p["loan"] and p["loan"] != it["f"]["acct"]:
            w += f"; NOTE the system suggested loan {p['loan']}"
        it["how"].append(f"{p['pid']}: {w}")


def free(t):
    return [p for p in led if not p["used"] and p["type"] == t]


# A. one posting, exactly the amount: tied to this loan, then held with this loan suggested, then the name, then the amount alone
#    (within 120 days of the value date), then tied to another loan of a customer of the same name
ranks = [
    (lambda p, f: p["state"].startswith("Tied") and p["loan"] == f["acct"], "tied to this loan by the system"),
    (lambda p, f: p["state"].startswith("Held") and p["loan"] == f["acct"], "held; the system's candidate loan"),
    (lambda p, f: p["state"].startswith("Held") and name_hit(p, f), "held; the customer's name is in the narration"),
    (lambda p, f: p["state"].startswith("Held") and days(p, f) <= 120, "held; same amount, near the value date"),
    (lambda p, f: p["state"].startswith("Tied") and (name_hit(p, f) or names.get(p["loan"], "").split()[0:1] == f["name"].split()[0:1]), "tied to another loan of the same customer"),
]
for test, why in ranks:
    for it in items:
        if it["posts"]:
            continue
        c = [p for p in free(it["t"]) if abs(p["amount"] - it["want"]) < 0.5 and test(p, it["f"])]
        c.sort(key=lambda p: days(p, it["f"]))
        if c:
            take(it, [c[0]], why)

# B. one held posting paying the fees of two loans together
for t in ("legal", "arrangement"):
    open_ = [it for it in items if it["t"] == t and not it["posts"]]
    for p in [p for p in free(t) if p["state"].startswith("Held")]:
        for a, b in combinations(open_, 2):
            if a["posts"] or b["posts"]:
                continue
            if abs(a["want"] + b["want"] - p["amount"]) < 0.5 and days(p, a["f"]) <= 60 and days(p, b["f"]) <= 60:
                p["used"] = True
                for it, other in ((a, b), (b, a)):
                    it["posts"].append({**p, "amount": it["want"], "shared": True})
                    it["how"].append(f"{p['pid']}: one posting of {p['amount']:,.2f} covers this loan and {other['f']['acct']} ({other['f']['name']})")
                break

# C. a posting plus its correction(s): up to three postings that sum to the amount, at least one tied to or naming the loan
for it in items:
    if it["posts"]:
        continue
    pool = [p for p in free(it["t"]) if (p["loan"] == it["f"]["acct"] or name_hit(p, it["f"]) or p["amount"] < 0
                                         or "over deducted" in (p["narration"] or "").lower()) and days(p, it["f"]) <= 250]
    done = False
    for k in (2, 3):
        for combo in combinations(pool, k):
            if abs(sum(p["amount"] for p in combo) - it["want"]) < 0.5 and any(p["loan"] == it["f"]["acct"] or name_hit(p, it["f"]) for p in combo):
                take(it, list(combo), "with the correcting entries")
                done = True
                break
        if done:
            break

# D. what is left: the nearest posting tied to or naming the loan, shown as a difference
for it in items:
    if it["posts"]:
        continue
    c = [p for p in free(it["t"]) if p["loan"] == it["f"]["acct"] or name_hit(p, it["f"])]
    c.sort(key=lambda p: abs(p["amount"] - it["want"]))
    if c:
        take(it, [c[0]], "nearest posting for this loan (amount differs)")

# E. a posting and its reversal that net to nil, attached to the loan the posting was tied to
for t in ("legal", "arrangement"):
    for a, b in combinations([p for p in free(t)], 2):
        if a["used"] or b["used"] or abs(a["amount"] + b["amount"]) > 0.5:
            continue
        loan = a["loan"] or b["loan"]
        it = next((x for x in items if x["t"] == t and x["f"]["acct"] == loan), None)
        if it:
            take(it, [a, b], "posted, then reversed (nets to nil)")

recon = []
for it in items:
    f, t = it["f"], it["t"]
    tied = sum(p["amount"] for p in it["posts"] if p["state"].startswith("Tied") and p["loan"] == f["acct"])
    ledger = round(sum(p["amount"] for p in it["posts"]), 2)
    diff = round(it["want"] - ledger, 2)
    if not it["posts"]:
        status = "In Finance's file, no posting found"
    elif abs(diff) > 0.5:
        status = "Amounts differ"
    elif all(p["state"].startswith("Tied") and p["loan"] == f["acct"] for p in it["posts"]):
        status = "Agrees"
    elif any(p["state"].startswith("Tied") and p["loan"] != f["acct"] for p in it["posts"]):
        status = "Agrees, but the system tied it to another loan"
    else:
        status = "Agrees once the held posting is tied to this loan"
    recon.append({"acct": f["acct"], "name": f["name"], "vd": f["value_date"], "type": t, "fin": it["want"], "tied": tied,
                  "held": ledger - tied, "ledger": ledger, "diff": diff, "status": status,
                  "postings": ", ".join(p["pid"] for p in it["posts"]), "dates": ", ".join(p["d"].isoformat() for p in it["posts"]),
                  "how": "; ".join(it["how"])})

left = [p for p in led if not p["used"]]

# ---------------------------------------------------------------- workbook
TEAL = "1F5C66"; HDR = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor=TEAL)
RED = PatternFill("solid", fgColor="FDE2E1"); AMB = PatternFill("solid", fgColor="FFF4CC"); GRN = PatternFill("solid", fgColor="E2F0D9")
WRAP = Alignment(wrap_text=True, vertical="top"); NUM = "#,##0.00"; LINE = Border(bottom=Side(style="thin", color="DDDDDD"))
wb = Workbook()


def table(ws, headers, data, widths, numcols=(), start=1):
    for c, h in enumerate(headers, 1):
        x = ws.cell(row=start, column=c, value=h); x.font = HDR; x.fill = HF; x.alignment = WRAP
    for i, row in enumerate(data, start + 1):
        for c, v in enumerate(row, 1):
            x = ws.cell(row=i, column=c, value=v); x.alignment = WRAP; x.border = LINE
            if headers[c - 1] in numcols: x.number_format = NUM
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[start].height = 32
    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(data)}"


cnt = defaultdict(int)
for r in recon: cnt[r["status"]] += 1
ws = wb.active; ws.title = "Summary"
lines = [
    ("MAIIC: 2025 fees, Finance's workbook against the general ledger", "t"),
    ("Prepared by Dupleix Institute, 15 October 2026. Finance's file: 'MAIIC EIR Assessment as of 31st December 2025 with account numbers.xlsx' (Tamanda, 15 September 2026, 34 loans). Ledger: every 2025 posting on GL 4871 (legal fees) and GL 4873 (arrangement fees) in the E-Banker extracts loaded in the IFRS 9 system. Nothing in either has been changed.", None),
    ("", None),
    ("1. Totals", "h"),
]
for i, (t, k) in enumerate(lines, 1):
    x = ws.cell(row=i, column=1, value=t); x.alignment = WRAP
    if k == "t": x.font = Font(bold=True, size=14, color=TEAL)
    if k == "h": x.font = Font(bold=True, color=TEAL)
r0 = len(lines) + 1
hd = ["", "Finance's file (sum of 34 loans)", "General ledger 2025 postings", "Difference"]
for c, h in enumerate(hd, 1):
    x = ws.cell(row=r0, column=c, value=h); x.font = HDR; x.fill = HF; x.alignment = WRAP
for j, (lab, t) in enumerate((("Legal fees, GL 4871", "legal"), ("Arrangement fees, GL 4873", "arrangement")), r0 + 1):
    for c, v in enumerate([lab, fin_tot[t], led_tot[t], fin_tot[t] - led_tot[t]], 1):
        x = ws.cell(row=j, column=c, value=v)
        if c > 1: x.number_format = NUM
r1 = r0 + 4
ws.cell(row=r1, column=1, value="2. Loan by loan").font = Font(bold=True, color=TEAL)
for c, h in enumerate(["Result", "Loan and fee lines"], 1):
    x = ws.cell(row=r1 + 1, column=c, value=h); x.font = HDR; x.fill = HF
order = ["Agrees", "Agrees once the held posting is tied to this loan", "Agrees, but the system tied it to another loan", "Amounts differ", "In Finance's file, no posting found"]
for j, s in enumerate(order, r1 + 2):
    ws.cell(row=j, column=1, value=s); ws.cell(row=j, column=2, value=cnt.get(s, 0))
r2 = r1 + 2 + len(order) + 1
ws.cell(row=r2, column=1, value=f"Ledger postings not matched to any loan in Finance's file: {len(left)} ({sum(p['amount'] for p in left):,.2f})").font = Font(bold=True)
notes = [
    ("3. What this means", "h"),
    ("The system ties a fee posting to a loan only when the loan's own account is a leg of the same batch. Fees credited by transfer (the 'By Trf' entries) have no loan leg, so the system holds them. Finance's file supplies the loan for those postings: where a held posting has exactly the amount Finance shows for a loan, the tie is listed on the 'By loan' sheet with the reason (the system's candidate loan, the customer's name in the narration, or the same amount only).", None),
    ("A tie proposed here is a suggestion for Finance to confirm. Once confirmed, each held posting is tied to its loan in the system (EIR > Fee Classification, held GL fees, 'Tie to loan'), which is recorded and audited. The workbook itself changes nothing.", None),
    ("Rows marked 'Amounts differ' or 'no posting found' need Finance to say which figure is right.", None),
]
for j, (t, k) in enumerate(notes, r2 + 2):
    x = ws.cell(row=j, column=1, value=t); x.alignment = WRAP
    if k == "h": x.font = Font(bold=True, color=TEAL)
ws.column_dimensions["A"].width = 95
for col in "BCD": ws.column_dimensions[col].width = 24

ws = wb.create_sheet("By loan")
H = ["Loan account", "Customer", "Value date", "Fee", "Finance's file (MWK)", "Ledger: tied by the system (MWK)",
     "Ledger: held posting matched here (MWK)", "Ledger total (MWK)", "Difference (MWK)", "Result", "Posting IDs", "Posting dates",
     "Why the held posting was matched", "Finance: confirm / correct", "Comment"]
data = [[r["acct"], r["name"], r["vd"], r["type"].title(), r["fin"], r["tied"], r["held"], r["ledger"], r["diff"], r["status"],
         r["postings"], r["dates"], r["how"], None, None] for r in sorted(recon, key=lambda r: (order.index(r["status"]) * -1, r["acct"], r["type"]))]
table(ws, H, data, [15, 28, 11, 11, 15, 15, 15, 15, 14, 26, 12, 22, 34, 18, 28],
      numcols=("Finance's file (MWK)", "Ledger: tied by the system (MWK)", "Ledger: held posting matched here (MWK)", "Ledger total (MWK)", "Difference (MWK)"))
for i, row in enumerate(data, 2):
    ws.cell(row=i, column=3).number_format = "yyyy-mm-dd"
    ws.cell(row=i, column=10).fill = GRN if row[9] == "Agrees" else AMB if row[9].startswith("Agrees") else RED
    for c in (14, 15): ws.cell(row=i, column=c).fill = AMB
    if row[0] == "106050000002" and row[3] == "Legal":
        for c in range(1, 16): ws.cell(row=i, column=c).fill = PatternFill("solid", fgColor="F8CBAD")
        ws.cell(row=i, column=15, value="QUESTION FOR TAMANDA (Q1): see the 'Questions for Tamanda' sheet")
        ws.cell(row=i, column=15).font = Font(bold=True)

# ---------------------------------------------------------------- questions to discuss with Tamanda
qs = wb.create_sheet("Questions for Tamanda", 0)
qs.cell(row=1, column=1, value="Questions to discuss with Tamanda").font = Font(bold=True, size=14, color=TEAL)
qs.cell(row=2, column=1, value="From the reconciliation of Finance's 2025 fee workbook to the general ledger, 15 October 2026.").alignment = WRAP
QH = ["#", "Loan", "Question", "Finance's workbook (MWK)", "General ledger (MWK)", "Loan system (MWK)", "Why it matters", "Tamanda's answer"]
chik = next(p for p in led if p["loan"] == "106050000002" and p["type"] == "legal")
questions = [[
    "Q1", "106050000002, Chikwendeni General Dealers (MWK 118,000,000, drawn 17 November 2025)",
    "The legal fee is 1,808,000 in your workbook but 1,808,800 on GL 4871 (posted " + chik["d"].isoformat() + ", " + (chik["ref"] or "") + "). "
    "Which figure is the fee charged to the customer? If 1,808,000 is right, the extra 800 on GL 4871 needs a correcting entry; if 1,808,800 is right, the workbook needs correcting. "
    "The loan system's charges table shows a third figure, 4,830,000 (and 1,949,800 for the arrangement fee, where your workbook and the GL agree on 3,540,000): is that table out of date for this loan?",
    1808000, 1808800, 4830000,
    "It is the only line in the 2025 workbook that does not agree with the GL to the kwacha (the 800 is the whole 2025 difference on legal fees). The system accepts it within the 1 percent band, but the fee that enters the loan's EIR should be the one charged.",
    None]]
table(qs, QH, questions, [5, 26, 70, 15, 15, 15, 45, 40], numcols=("Finance's workbook (MWK)", "General ledger (MWK)", "Loan system (MWK)"), start=4)
qs.freeze_panes = None
for c in range(1, 9):
    qs.cell(row=5, column=c).fill = PatternFill("solid", fgColor="F8CBAD")
qs.cell(row=5, column=8).fill = AMB
qs.row_dimensions[5].height = 150

ws = wb.create_sheet("Ledger postings")
H2 = ["Posting ID", "State in the system", "GL", "Fee", "Date", "Amount (MWK)", "Narration", "Batch / reference",
      "Loan the system holds or suggests", "Matched to loan here"]
match_of = {}
for r in recon:
    for pid in r["postings"].split(", "):
        match_of.setdefault(pid, "")
        if pid: match_of[pid] = (match_of[pid] + ", " if match_of[pid] else "") + r["acct"]
data2 = [[p["pid"], p["state"], p["gl"], p["type"].title(), p["d"], p["amount"], p["narration"], p["ref"], p["loan"], match_of.get(p["pid"], "NOT MATCHED")]
         for p in sorted(led, key=lambda p: (p["d"], p["gl"]))]
table(ws, H2, data2, [9, 24, 7, 11, 11, 15, 50, 22, 16, 16], numcols=("Amount (MWK)",))
for i, row in enumerate(data2, 2):
    ws.cell(row=i, column=5).number_format = "yyyy-mm-dd"
    if row[9] == "NOT MATCHED": ws.cell(row=i, column=10).fill = RED

wb.save(OUT)
print("finance", fin_tot, "ledger", led_tot, dict(cnt), "unmatched postings", len(left), round(sum(p["amount"] for p in left), 2))
for r in recon:
    if r["status"] != "Agrees":
        print(r["acct"], r["name"][:25], r["type"], r["fin"], r["ledger"], r["diff"], r["status"], r["how"])
for p in left:
    print("LEFT", p["pid"], p["gl"], p["d"], p["amount"], (p["narration"] or "")[:60], p["loan"])
