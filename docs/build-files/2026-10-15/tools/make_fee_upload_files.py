"""Pre-filled fee workbook uploads for Finance (15 Oct 2026): the 2025 file ready to load through EIR Data Intake (type fees),
and October-December 2024 to complete. Each is run through the system's gate (FeeWorkbookTieOut, read only) first."""
import csv, json, subprocess, sys, datetime, os
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTDIR = sys.argv[1]
S = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP = r"C:\xampp\htdocs\MAICC-IFRS9"
FIN25 = r"C:\Users\DELL\OneDrive\2026\Projects\MAIIC\2. Documents from clients\EIR Assesment\MAIIC EIR Assessment as of 31st December 2025 with account numbers.xlsx"
FIN24 = r"C:\Users\DELL\OneDrive\2026\Projects\MAIIC\2. Documents from clients\EIR Assesment\MAIIC EIR Assessment as of 31st December 2025.xlsx"
COLS = ["contract_id", "fee_type", "amount", "transaction_date", "gl_account_ref", "cashflow_direction", "currency", "description",
        "source_system", "source_reference", "external_transaction_id"]
GL = {"legal": "4871", "arrangement": "4873"}


def gate(rows, tag):
    a, b = os.path.join(S, f"gate_{tag}_in.json"), os.path.join(S, f"gate_{tag}_out.json")
    json.dump(rows, open(a, "w"))
    subprocess.run(["php", os.path.join(S, "tieout_json.php"), a, b], cwd=APP, check=True)
    return json.load(open(b))


# ---------------------------------------------------------------- 2025: Finance's file with account numbers
ws = openpyxl.load_workbook(FIN25, data_only=True).active
rows25, meta25 = [], []
for i, r in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
    if not r[3] or not str(r[3]).strip().isdigit():
        continue
    vd = r[7] if not isinstance(r[7], str) else datetime.datetime.strptime(r[7].strip(), "%d/%m/%Y")
    for t, col in (("legal", 18), ("arrangement", 20)):
        if not r[col]:
            continue
        rows25.append({"contract_id": str(r[3]).strip(), "fee_type": t, "amount": round(float(r[col]), 2), "transaction_date": vd.date().isoformat(),
                       "gl_account_ref": GL[t], "cashflow_direction": "RECEIVED", "currency": "MWK",
                       "description": f"{t.title()} fee, {str(r[5] or r[4]).strip()} (Finance EIR assessment 31 Dec 2025)",
                       "source_system": "MAIIC Finance EIR assessment 2025",
                       "source_reference": f"MAIIC EIR Assessment as of 31st December 2025 with account numbers.xlsx, row {i}",
                       "external_transaction_id": f"FIN-EIR2025-R{i}-{t.upper()}"})
        meta25.append({"customer": str(r[5] or r[4]).strip(), "row": i})
g25 = gate(rows25, "2025")
notes25 = {}
for l in g25["lines"]:
    notes25[l["row"]] = ("; ".join(l["how"]), l.get("difference", 0))
conf25 = {c["row"]: c["note"] for c in g25["conflicts"]}
tie25 = {t["row"]: t["posting"] for t in g25["ties"]}
csv25 = os.path.join(OUTDIR, "MAIIC fee workbook upload - 2025 - ready - 15 Oct 2026.csv")
with open(csv25, "w", newline="", encoding="utf-8") as fh:
    w = csv.DictWriter(fh, fieldnames=COLS)
    w.writeheader()
    w.writerows(rows25)

# ---------------------------------------------------------------- Oct-Dec 2024: Finance's lines of loans drawn in the E-Banker period
ws = openpyxl.load_workbook(FIN24, data_only=True)["EIR Assessment December 2025"]
ACCT = {  # name in Finance's file -> (account, basis)
    "Agwenda Investment Limited": ("104430000064", "tied by the system (the loan in the batch)"),
    "Eye Care Centre Limited": ("104460000105", "the system's candidate loan (name and drawdown date)"),
    "Microloan Foundation Officer": ("104430000069", "tied by the system (the loan in the batch)"),
    "Chimwemwe Mafeni Transport": ("104460000104", "tied by the system (the loan in the batch)"),
    "NyamNyam Processors Limited": ("104430000068", "TO CONFIRM: 'Nyamunyamu processors limited', in the loan book from November 2024 (MWK 200m); 104430000067 'Nyamnyam processors limited' was opened 20 Nov 2024"),
    "Premium Foods Limited": ("104460000107", "TO CONFIRM: 'Premium meat suppliers', opened 31 Dec 2024 (MWK 75m); 105810000002 of the same name was opened 19 Dec 2024"),
    "Kwalingana Investments Limited": ("", "NEEDED: no E-Banker account under this name"),
    "Napoleon dzombe T/A Santhe Business Park": ("", "NEEDED: no E-Banker account under this name"),
}
rows24, meta24 = [], []
for i, r in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
    name, vd = (str(r[4]).strip() if r[4] else None), r[6]
    if not name or not isinstance(vd, datetime.datetime) or vd.date() < datetime.date(2024, 10, 1) or vd.year > 2024:
        continue
    acct, basis = ACCT.get(name, ("", "NEEDED"))
    for t, col in (("legal", 17), ("arrangement", 19)):
        if not r[col]:
            continue
        rows24.append({"contract_id": acct, "fee_type": t, "amount": round(float(r[col]), 2), "transaction_date": vd.date().isoformat(),
                       "gl_account_ref": GL[t], "cashflow_direction": "RECEIVED", "currency": "MWK",
                       "description": f"{t.title()} fee, {name} (Finance EIR assessment, all loans tested)",
                       "source_system": "MAIIC Finance EIR assessment 2025", "source_reference": f"MAIIC EIR Assessment as of 31st December 2025.xlsx, row {i}",
                       "external_transaction_id": f"FIN-EIR2024-R{i}-{t.upper()}"})
        meta24.append({"customer": name, "basis": basis})
g24 = gate(rows24, "2024q4")

# ---------------------------------------------------------------- the review workbook
TEAL = "1F5C66"; HDR = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor=TEAL)
RED = PatternFill("solid", fgColor="FDE2E1"); AMB = PatternFill("solid", fgColor="FFF4CC"); GRN = PatternFill("solid", fgColor="E2F0D9")
Q = PatternFill("solid", fgColor="F8CBAD")
WRAP = Alignment(wrap_text=True, vertical="top"); NUM = "#,##0.00"; LINE = Border(bottom=Side(style="thin", color="DDDDDD"))


def table(ws, headers, data, widths, numcols=(), start=1, inputs=()):
    for c, h in enumerate(headers, 1):
        x = ws.cell(row=start, column=c, value=h); x.font = HDR if h not in inputs else Font(bold=True); x.fill = HF if h not in inputs else AMB; x.alignment = WRAP
    for i, row in enumerate(data, start + 1):
        for c, v in enumerate(row, 1):
            x = ws.cell(row=i, column=c, value=v); x.alignment = WRAP; x.border = LINE
            if headers[c - 1] in numcols: x.number_format = NUM
            if headers[c - 1] in inputs: x.fill = AMB
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[start].height = 32
    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(data)}"


wb = Workbook()
ws = wb.active; ws.title = "Read me"
txt = [
    ("MAIIC fee workbook: the files to upload", "t"),
    ("Prepared by Dupleix Institute for MAIIC Finance, 15 October 2026. The IFRS 9 system now loads Finance's fee workbook only when it ties to the GL fee postings, line by line and month by month. These are Finance's own figures, put in the layout the system's fee import reads.", None),
    ("", None),
    ("1. 2025: ready to load", "h"),
    (f"File: '{os.path.basename(csv25)}' ({len(rows25)} lines, from 'MAIIC EIR Assessment as of 31st December 2025 with account numbers.xlsx'). Checked against the system's gate today: {g25['status']}. "
     f"On loading, {len(g25['ties'])} GL postings the system holds without a loan are tied to their loans. Please read the '2025 lines' sheet before loading, in particular the matches made on amount and date only, the two Mthirakuwiri postings and Q1 (Chikwendeni).", None),
    ("How to load: EIR & Revenue Recognition > EIR Data > EIR Data Intake (or Fee Classification > Import fees), type 'Fees and transaction costs', choose the CSV, check the column mapping (each column has the system's own name), Check file, then Import. If the gate refuses the file, nothing loads and the exception file in the import history says why.", None),
    ("", None),
    ("2. October to December 2024: to complete before loading", "h"),
    (f"Finance's lines for the {len({m['customer'] for m in meta24})} loans drawn from October 2024 (from 'MAIIC EIR Assessment as of 31st December 2025.xlsx', which has no account numbers). Gate today: {g24['status']}, for two reasons, both on the '2024 Oct-Dec' sheet: "
     "four loans need their E-Banker account confirmed or given, and November 2024 has GL fee postings that are not in Finance's file (other loans' fees and reallocation journals). The system needs every fee posting of a month in the file, so please say whose each one is (or that it is not a loan fee).", None),
    ("", None),
    ("3. Loans drawn before October 2024", "h"),
    ("Their fees are in the take-on workbook MAIIC completed on 13 October 2026 and stay there; the system refuses a line that repeats one. E-Banker holds no fee posting before 3 October 2024.", None),
]
for i, (t, k) in enumerate(txt, 1):
    x = ws.cell(row=i, column=1, value=t); x.alignment = WRAP
    if k == "t": x.font = Font(bold=True, size=14, color=TEAL)
    if k == "h": x.font = Font(bold=True, color=TEAL)
ws.column_dimensions["A"].width = 120

ws = wb.create_sheet("2025 lines")
H = ["CSV row", "Loan", "Customer", "Fee", "Amount (MWK)", "Date", "How the system matched it", "Difference (MWK)", "Held posting tied on loading", "To look at", "Finance: OK to load?", "Comment"]
data = []
for k, (row, m) in enumerate(zip(rows25, meta25), 2):
    how, diff = notes25.get(k, ("", 0))
    look = []
    if "no name" in how: look.append("Matched on amount and date only: please confirm the loan")
    if k in conf25: look.append(conf25[k])
    if row["contract_id"].lstrip("0") == "106050000002" and row["fee_type"] == "legal": look.append("Q1 FOR TAMANDA: 1,808,000 here, 1,808,800 on GL 4871, 4,830,000 in the loan system")
    data.append([k, row["contract_id"], m["customer"], row["fee_type"].title(), row["amount"], row["transaction_date"], how, diff or None, tie25.get(k), "; ".join(look) or None, None, None])
table(ws, H, data, [7, 16, 28, 11, 15, 11, 50, 12, 12, 50, 12, 28], numcols=("Amount (MWK)", "Difference (MWK)"), inputs=("Finance: OK to load?", "Comment"))
for i, row in enumerate(data, 2):
    if row[9]:
        ws.cell(row=i, column=10).fill = Q if "Q1" in row[9] else RED

ws = wb.create_sheet("2024 Oct-Dec")
H = ["File row", "Loan (E-Banker account)", "Customer", "Fee", "Amount (MWK)", "Value date", "Where the account comes from", "Gate today", "Finance: account", "Comment"]
ref24 = {r["row"]: r["reason"] for r in g24["refusals"] if r["row"]}
how24 = {l["row"]: "; ".join(l["how"]) for l in g24["lines"]}
data = []
for k, (row, m) in enumerate(zip(rows24, meta24), 2):
    data.append([k, row["contract_id"] or None, m["customer"], row["fee_type"].title(), row["amount"], row["transaction_date"], m["basis"],
                 ref24.get(k) or how24.get(k) or "", None, None])
table(ws, H, data, [7, 18, 32, 11, 15, 11, 50, 50, 18, 28], numcols=("Amount (MWK)",), inputs=("Finance: account", "Comment"))
for i, row in enumerate(data, 2):
    ws.cell(row=i, column=7).fill = GRN if not ("NEEDED" in row[6] or "CONFIRM" in row[6]) else (RED if "NEEDED" in row[6] else AMB)
r0 = len(data) + 4
ws.cell(row=r0, column=1, value="GL fee postings of October to December 2024 that are in none of Finance's lines (the system needs each one accounted for)").font = Font(bold=True, color=TEAL)
H2 = ["Posting", "Date", "GL", "Amount (MWK)", "Narration", "Loan the system has or the name suggests", "Finance: whose fee is it (loan), or why it is not a loan fee"]
NAMEHINT = {"rose harris": "104450000047 Rose Harris Investment (opened 2022)", "mawindo": "104450000026 Mawindo Enterprise (opened 2022)",
            "kaning": "104450000048 Kaning'a General Dealer (take-on loan)", "chikwendeni": "104450000028 Chikwendeni General Dealers (opened 2022)",
            "santhe": "Santhe (no E-Banker account found)", "kwalingana": "Kwalingana (no E-Banker account found)", "nyamnyam": "104430000068 or 104430000067 (NyamNyam)"}
lst = []
for u in g24["unaccounted"]:
    for p in u["postings"]:
        n = (p["narration"] or "").lower()
        hint = p["loan"] or next((v for k2, v in NAMEHINT.items() if k2 in n), "")
        lst.append([p["posting"], p["date"], u["gl"], p["amount"], p["narration"], hint, None])
for c, h in enumerate(H2, 1):
    x = ws.cell(row=r0 + 1, column=c, value=h); x.font = HDR if c < 7 else Font(bold=True); x.fill = HF if c < 7 else AMB; x.alignment = WRAP
for i, row in enumerate(lst, r0 + 2):
    for c, v in enumerate(row, 1):
        x = ws.cell(row=i, column=c, value=v); x.alignment = WRAP; x.border = LINE
        if c == 4: x.number_format = NUM
        if c == 7: x.fill = AMB
ws.cell(row=r0 + 2 + len(lst) + 1, column=1, value="A posting and its reversal (for example the 84,692,000 mis-posting of 30 November 2024 and its reversal) cancel and are not listed.").alignment = WRAP

out = os.path.join(OUTDIR, "MAIIC fee workbook uploads - for Finance - 15 Oct 2026.xlsx")
wb.save(out)
print("2025", g25["status"], len(rows25), "lines, ties", len(g25["ties"]), "conflicts", len(g25["conflicts"]))
print("2024q4", g24["status"], len(rows24), "lines; refusals", len(g24["refusals"]), "; unaccounted postings", sum(len(u["postings"]) for u in g24["unaccounted"]))
print(out); print(csv25)
