import openpyxl, sys
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
out = openpyxl.Workbook(); out.remove(out.active)
for name in ["ASS1-Internal Grades", "PDs-S&P"]:
    ws = wb[name]; o = out.create_sheet(name); n = 0
    for row in ws.iter_rows(values_only=True):
        o.append(list(row)); n += 1
    print(name, n, "rows")
out.save(dst)
