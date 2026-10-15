import openpyxl, sys, collections
wb = openpyxl.load_workbook(sys.argv[1])
for ws in wb: print(ws.title, ws.max_row, ws.max_column)
ws = wb["Fee lines"]; h = [c.value for c in ws[1]]
rows = [dict(zip(h, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
print("blank customer", sum(1 for r in rows if not r["Customer"]))
print(collections.Counter((r["Rule that would apply"] or "NONE")[:60] for r in rows if r["Needs a decision"] == "Yes").most_common(12))
fl = collections.Counter()
for r in rows:
    for x in (r["Flags to check"] or "").split("; "):
        if x: fl[x.split(":")[0][:70]] += 1
print(fl)
print([ (r["Loan account"], r["Description"], r["Flags to check"]) for r in rows if (r["Rule that would apply"] or "") == ""][:3])
ws = wb["Rules to approve"]; print([(c[0].value, c[7].value) for c in ws.iter_rows(min_row=2) if c[7].value][:10])
