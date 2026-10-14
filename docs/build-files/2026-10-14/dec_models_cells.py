import openpyxl, os, glob, warnings
warnings.filterwarnings("ignore")
base = r"C:\Users\wadza\OneDrive\2026\Projects\MAIIC\2. Documents from clients\New Doc Received 13 Oct"
files = sorted(f for f in glob.glob(os.path.join(base, "*.xlsx")) if "ECL" in f)
WANT = {
    "ECLComp": ["P2:S7", "AA2:AB5", "CA4:CB7", "N10:N22", "AX7"],
    "ECLComp(InternalPDByCount) ": ["L2:N6", "L10:N22"],
    "ASS1a Internal Grades-Indv Asse": ["B5:D12"],
    "PortfolioReport": ["AD4:AN4", "AO4:AS5", "AD5"],
    "Loan Book": ["Y3:AE5"],
    "Checklist": ["L19:M24"],
    "Sheet2": ["B2:C19"],
}
def show(wb, wv, sheet, rng):
    if sheet not in wb.sheetnames:
        print(f"  [{sheet}] missing"); return
    ws = wb[sheet]; vs = wv[sheet]
    print(f"  [{sheet}] {rng}")
    cells = ws[rng]
    if not isinstance(cells, tuple):
        cells = ((cells,),)
    for row in cells:
        for c in row:
            if c.value is None: continue
            v = c.value; vv = vs[c.coordinate].value
            s = repr(v)[:160]
            if isinstance(v, str) and v.startswith("="):
                s += f"  -> {vv!r}"
            print(f"     {c.coordinate}: {s}")
for f in files:
    print("=" * 90); print(os.path.basename(f))
    wb = openpyxl.load_workbook(f); wv = openpyxl.load_workbook(f, data_only=True)
    for sheet, rngs in WANT.items():
        for r in rngs: show(wb, wv, sheet, r)
    # stage counts and bracket counts in ECLComp
    if "ECLComp" in wv.sheetnames:
        vs = wv["ECLComp"]; from collections import Counter
        br = Counter(); st = Counter(); rbm = Counter()
        for r in range(11, vs.max_row + 1):
            p = vs.cell(r, 16).value; s = vs.cell(r, 18).value; n = vs.cell(r, 14).value; e = vs.cell(r, 5).value
            if e is None and p is None: continue
            br[p] += 1; st[s] += 1; rbm[n] += 1
        print("  ECLComp brackets:", dict(br)); print("  ECLComp stages:", dict(st)); print("  ECLComp RBM CLASS values:", dict(rbm))
