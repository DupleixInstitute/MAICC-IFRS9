import openpyxl, os, sys, warnings
warnings.filterwarnings("ignore")
base = r"C:\Users\wadza\OneDrive\2026\Projects\MAIIC\2. Documents from clients\New Doc Received 13 Oct"
MAIN = os.path.join(base, "MAIIC & FInES ECL Model - 14th December 2025_Final.xlsx")
MEGA = os.path.join(base, "Mega Farms ECL Model - 14th December 2025_Final.xlsx")

def dump(ws, rng, label, both=None):
    print(f"\n#### {ws.title} {rng}  ({label})")
    for row in ws[rng]:
        for c in row:
            v = c.value
            if v is None:
                continue
            s = repr(v)
            if len(s) > 220:
                s = s[:220] + "..."
            extra = ""
            if both is not None:
                vv = both[ws.title][c.coordinate].value
                if isinstance(v, str) and v.startswith("=") and vv is not None:
                    extra = f"   -> {vv!r}"
            print(f"  {c.coordinate}: {s}{extra}")

for path, tag in [(MAIN, "MAIN"), (MEGA, "MEGA")]:
    print("\n" + "=" * 100 + f"\n{tag}: {os.path.basename(path)}")
    wb = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames
    print("sheets:", names)
    # ECLComp header block and first 3 data rows
    if "ECLComp" in names:
        ws = wb["ECLComp"]
        dump(ws, "A1:CU12", "ECLComp header rows", wv)
        dump(ws, "A13:CU15", "ECLComp first 3 facility rows (formulas -> values)", wv)
    for nm in names:
        if nm.startswith("ECLComp(InternalPDByCount"):
            dump(wb[nm], "L1:Z12", "InternalPDByCount header incl RBM Asset Classification", wv)
            dump(wb[nm], "L13:Z14", "InternalPDByCount first rows", wv)
    for nm in ["ASS1a Internal Grades-Indv Asse", "ASS1-Internal Grades", "Validation Tables", "Checklist", "Sheet2", "Closed Accounts"]:
        if nm in names:
            ws = wb[nm]
            dump(ws, f"A1:{openpyxl.utils.get_column_letter(min(ws.max_column,16))}{min(ws.max_row,62)}", nm, wv)
    if "ASS4-Security" in names:
        dump(wb["ASS4-Security"], "B1:O42", "ASS4-Security", wv)
    if "PortfolioReport" in names:
        ws = wb["PortfolioReport"]
        dump(ws, "A1:AX6", "PortfolioReport header + first data row", wv)
        dump(ws, "A7:AX8", "PortfolioReport rows 7-8", wv)
    if "Loan Book" in names:
        ws = wb["Loan Book"]
        dump(ws, "A1:AT7", "Loan Book header + first rows", wv)
        dump(ws, "S125:AT140", "Loan Book totals / NPL block", wv)
    if "KPIs" in names:
        dump(wb["KPIs"], "B2:G13", "KPIs", wv)
    if "Summary(OriginalMethod)" in names:
        dump(wb["Summary(OriginalMethod)"], "B1:O28", "Summary", wv)
