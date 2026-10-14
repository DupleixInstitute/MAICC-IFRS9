from pathlib import Path
F = Path(r"C:\xampp\htdocs\MAICC-IFRS9\docs\build-files\spec_render.py")
s = F.read_text(encoding="utf-8")
def sub(a, b):
    global s
    assert a in s, a[:80]
    s = s.replace(a, b, 1)

sub("and justify (yes: paragraphs and list items fully justified);",
    "and justify (yes: paragraphs and list items fully justified); toc (yes: a table of contents with page\nnumbers after the cover block: built in two passes in the PDF, a TOC field Word fills in on opening);")

# ---------------- docx: TOC field after the cover block, and ask Word to update fields on opening
sub("""    for kind, content in blocks:
        if kind == "h2": doc.add_paragraph(content, style="Heading 2")""",
    """    if meta.get("toc", "").lower() in ("yes", "true", "1"):
        h = doc.add_paragraph(); r = h.add_run("Contents"); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY
        p = doc.add_paragraph(); run = p.add_run()
        for tag, text in (("begin", None), (None, 'TOC \\\\o "2-3" \\\\h \\\\z \\\\u'), ("separate", None), (None, None), ("end", None)):
            if tag:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
            elif text:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)
            else:
                t = OxmlElement("w:t"); t.text = "Right-click and choose Update Field if the page numbers do not show."; run._r.append(t)
        upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); doc.settings.element.append(upd)
        from docx.enum.text import WD_BREAK
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for kind, content in blocks:
        if kind == "h2": doc.add_paragraph(content, style="Heading 2")""")

# ---------------- pdf: TableOfContents with multiBuild
sub("""    from reportlab.lib.enums import TA_JUSTIFY""",
    """    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.platypus.tableofcontents import TableOfContents""")
sub("""    s.append(Spacer(1, 8))
    for kind, content in blocks:""",
    """    s.append(Spacer(1, 8))
    TOC = meta.get("toc", "").lower() in ("yes", "true", "1")
    if TOC:
        toc = TableOfContents(dotsMinLevel=0)
        toc.levelStyles = [ParagraphStyle("t0", fontName="Seg-B", fontSize=10, leading=15, leftIndent=0, firstLineIndent=0, textColor=NAVY, spaceBefore=3),
                           ParagraphStyle("t1", fontName="Seg", fontSize=9.2, leading=13, leftIndent=14, firstLineIndent=0, textColor=GREY)]
        s.append(Paragraph("Contents", S["h2"])); s.append(toc); s.append(PageBreak())
    for kind, content in blocks:""")
sub("""    doc = BaseDocTemplate(str(out), pagesize=A4,""",
    """    class _Doc(BaseDocTemplate):
        def afterFlowable(self, f):
            if TOC and isinstance(f, Paragraph) and f.style.name in ("h2", "h3") and f.getPlainText() != "Contents":
                self.notify("TOCEntry", (0 if f.style.name == "h2" else 1, f.getPlainText(), self.page))
    doc = _Doc(str(out), pagesize=A4,""")
sub("""    doc.build(s)""", """    doc.multiBuild(s) if TOC else doc.build(s)""")
F.write_text(s, encoding="utf-8")
print("toc patched")
