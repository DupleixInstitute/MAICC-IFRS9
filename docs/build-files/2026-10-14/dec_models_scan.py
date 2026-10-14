import openpyxl, re, sys, glob, os
base = r"C:\Users\wadza\OneDrive\2026\Projects\MAIIC\2. Documents from clients\New Doc Received 13 Oct"
files = [f for f in glob.glob(os.path.join(base, "*.xlsx")) if "ECL" in f]
KEYS = re.compile(r"(past ?due|days|dpd|classif|provision|suspense|restructur|reschedul|special mention|substandard|doubtful|loss|standard|performing|npl|directive|rbm|prudential|regulat|bracket|bucket|stage|sicr|cure|moratori|arrear)", re.I)
for f in files:
    print("=" * 100)
    print(os.path.basename(f))
    wb = openpyxl.load_workbook(f, data_only=False, read_only=False)
    for ws in wb.worksheets:
        print(f"--- sheet [{ws.title}] dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
    # keyword hits per sheet (limit rows scanned to 400 and cols to 60)
    for ws in wb.worksheets:
        hits = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400), max_col=min(ws.max_column, 60)):
            for c in row:
                v = c.value
                if isinstance(v, str) and KEYS.search(v) and len(v) < 200:
                    hits.append(f"{c.coordinate}={v!r}")
        if hits:
            print(f"## [{ws.title}] {len(hits)} hits")
            for h in hits[:80]:
                print("   ", h)
